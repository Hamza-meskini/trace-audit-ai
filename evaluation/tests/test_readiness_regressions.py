"""Production boundary regressions; no live credentials or remote services."""
import asyncio
import importlib
import io
from unittest.mock import AsyncMock, MagicMock, patch

import httpx
import pytest
from fastapi import HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker

from app import database
from app.config import settings
from app.api.documents import upload_document, upload_documents
from app.api.projects import delete_project
from app.api.visitors import _csv_cell
from app.models.project import Project


def test_transactions_commit_rollback_and_cancel():
    async def run():
        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        async with engine.begin() as conn:
            await conn.run_sync(database.Base.metadata.create_all)
        try:
            with patch.object(database, "async_session", sessions), patch.object(settings, "DATABASE_BACKEND", "sqlite"):
                async with database.get_db_context() as db:
                    db.add(Project(id="saved", name="Saved", product_name="Unit", audit_id="A1"))
                for failure in (RuntimeError, asyncio.CancelledError):
                    with pytest.raises(failure):
                        async with database.get_db_context() as db:
                            db.add(Project(id="failed", name="Failed", product_name="Unit", audit_id="A2"))
                            await db.flush()
                            raise failure()
                async with sessions() as db:
                    assert list((await db.scalars(select(Project.id))).all()) == ["saved"]
        finally:
            await engine.dispose()
    asyncio.run(run())


def test_databricks_session_closes_when_commit_fails():
    async def run():
        sync = MagicMock()
        sync.commit.side_effect = RuntimeError("commit failed")
        with pytest.raises(RuntimeError):
            async with database.AsyncDatabricksSession(sync):
                pass
        sync.rollback.assert_called_once()
        sync.close.assert_called_once()
    asyncio.run(run())


def test_static_files_cannot_escape_build(tmp_path):
    import app.main as main
    build = tmp_path / "client"
    build.mkdir()
    (build / "index.html").write_text("<html><head></head><body>APP</body></html>")
    (tmp_path / "private.txt").write_text("PRIVATE")

    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url="http://test") as client:
            assert (await client.get("/projects")).status_code == 200
            identity = (await client.get("/api/me", headers={"X-Forwarded-User": "reviewer", "X-Forwarded-Email": "reviewer@example.test"})).json()
            assert identity["user"] == "reviewer"
            assert identity["provider"] == "databricks_sso"
            assert identity["authenticated"] is True
            for path in ("/%2e%2e/private.txt", "/..%5cprivate.txt", "/api/no-such-route"):
                response = await client.get(path)
                assert response.status_code == 404
                assert "PRIVATE" not in response.text
    try:
        with patch.object(settings, "FRONTEND_DIR", str(build)):
            importlib.reload(main)
            asyncio.run(run())
    finally:
        importlib.reload(main)


def test_upload_boundaries_and_failed_persistence(tmp_path):
    async def run():
        db = AsyncMock()
        db.add = MagicMock()
        result = MagicMock()
        result.scalar_one_or_none.return_value = object()
        db.execute.return_value = result
        with patch.object(settings, "UPLOAD_DIR", str(tmp_path)), patch("app.api.documents.MAX_UPLOAD_SIZE_BYTES", 4):
            for content, status in ((b"", 400), (b"12345", 413)):
                file = UploadFile(filename="sample.pdf", file=io.BytesIO(content))
                with pytest.raises(HTTPException) as error:
                    await upload_document("p1", file, "", "v1", db)
                assert error.value.status_code == status
            db.commit.side_effect = RuntimeError("database unavailable")
            with pytest.raises(RuntimeError):
                await upload_document("p1", UploadFile(filename="sample.pdf", file=io.BytesIO(b"1234")), "", "v1", db)
            assert not list(tmp_path.rglob("*.pdf"))
    asyncio.run(run())


def test_batch_upload_uses_one_commit(tmp_path):
    async def run():
        db = AsyncMock()
        db.add = MagicMock()
        result = MagicMock()
        result.scalar_one_or_none.return_value = object()
        db.execute.return_value = result
        files = [
            UploadFile(filename="requirements.pdf", file=io.BytesIO(b"1234")),
            UploadFile(filename="evidence.docx", file=io.BytesIO(b"abcd")),
        ]
        with patch.object(settings, "UPLOAD_DIR", str(tmp_path)):
            saved = await upload_documents("p1", files, "", "v1", db)
        assert [item.original_filename for item in saved] == ["requirements.pdf", "evidence.docx"]
        assert db.commit.await_count == 1
        assert db.refresh.await_count == 0
        stored = list(tmp_path.rglob("*.pdf")) + list(tmp_path.rglob("*.docx"))
        assert len(stored) == 2
    asyncio.run(run())


@pytest.mark.parametrize("root", ["", "/apps/traceaudit"])
def test_nested_pages_resolve_assets_from_app_root(tmp_path, root):
    import app.main as main
    from starlette.requests import Request
    index = tmp_path / "index.html"
    index.write_text('<html><head><link href="./assets/style.css"></head>'
                     '<body><script src="/./assets/app.js"></script></body></html>')
    try:
        with patch.object(settings, "FRONTEND_DIR", str(tmp_path)):
            importlib.reload(main)
            request = Request({"type": "http", "root_path": root})
            html = main._render_index(index, request).body.decode()
            assert f'<base href="{root}/">' in html
            if root:
                assert f'src="{root}/assets/app.js"' in html
                assert 'window.__DATABRICKS_ROOT_PATH__ = "/apps/traceaudit"' in html
    finally:
        importlib.reload(main)


@pytest.mark.parametrize("status", ["queued", "running", "cancelling"])
def test_active_project_cannot_be_deleted(status):
    async def run():
        db = AsyncMock()
        with patch("app.api.projects.audit_progress.latest", return_value={"status": status}):
            with pytest.raises(HTTPException) as error:
                await delete_project("p1", db)
            assert error.value.status_code == 409
            db.delete.assert_not_called()
    asyncio.run(run())


@pytest.mark.parametrize("value", ["=1+1", " +SUM(1,2)", "@SUM(1,2)", "-1+1", "\tformula"])
def test_csv_formulas_are_escaped(value):
    assert _csv_cell(value) == "'" + value
    assert _csv_cell("ordinary text") == "ordinary text"


def test_excel_audit_report_generation():
    import io
    import openpyxl
    from app.services.excel_export import generate_audit_report_excel
    from app.models.requirement import Requirement
    from app.models.finding import Finding

    async def run():
        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        async with engine.begin() as conn:
            await conn.run_sync(database.Base.metadata.create_all)
        try:
            async with sessions() as db:
                db.add(Project(
                    id="proj-xl-1",
                    name="Battery Unit Test",
                    product_name="BMS-X",
                    audit_id="TA-2026-TEST",
                    company="Acme Energy",
                ))
                db.add(Requirement(
                    id="req-xl-1",
                    project_id="proj-xl-1",
                    req_code="REQ-001",
                    title="Active Discharge",
                    description="The battery system shall discharge active voltage within 5 seconds.",
                    coverage_status="Supported",
                    severity="Critical",
                    confidence=0.98,
                    review_state="Approved",
                    ai_analysis="Full compliance verified via hardware telemetry log.",
                    ai_recommendation="Maintain monitoring in continuous test.",
                ))
                db.add(Finding(
                    id="find-xl-1",
                    project_id="proj-xl-1",
                    finding_code="FIND-001",
                    finding_type="Missing evidence",
                    severity="High",
                    review_state="Open",
                    description="Secondary temperature sensor trace not attached.",
                ))
                await db.commit()

                excel_bytes = await generate_audit_report_excel(db, "proj-xl-1")
                assert len(excel_bytes) > 0

                wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
                assert "Executive Summary" in wb.sheetnames
                assert "Requirement Register" in wb.sheetnames
                assert "Evidence Traceability" in wb.sheetnames
                assert "Findings & Issues" in wb.sheetnames

                ws_summary = wb["Executive Summary"]
                assert ws_summary["B4"].value == "Battery Unit Test"
                assert ws_summary["E4"].value == "TA-2026-TEST"

                ws_reqs = wb["Requirement Register"]
                assert ws_reqs.max_row >= 2
                assert ws_reqs["A2"].value == "REQ-001"
                assert ws_reqs["E2"].value == "Supported"
                assert "discharge active voltage" in ws_reqs["K2"].value
                assert "Full compliance verified" in ws_reqs["L2"].value

                ws_finds = wb["Findings & Issues"]
                assert ws_finds.max_row >= 2
                assert ws_finds["A2"].value == "FIND-001"
        finally:
            await engine.dispose()

    asyncio.run(run())
